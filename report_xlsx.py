"""Tạo file Excel báo cáo tuần / tháng (openpyxl). Dùng công thức cho các ô tổng và % thay đổi."""
from datetime import datetime, timedelta, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

VN = timezone(timedelta(hours=7))
BROWN, CREAM = "4B2E1E", "F6EFE4"
ST = {"pending": "Chờ xác nhận", "preparing": "Đang pha chế", "ready": "Sẵn sàng", "completed": "Hoàn tất", "cancelled": "Đã hủy"}
TY = {"dine_in": "Tại quán", "takeaway": "Mang đi", "delivery": "Giao hàng"}
PAY = {"cash": "Tiền mặt", "transfer": "Chuyển khoản", "momo": "MoMo"}

F = lambda **k: Font(name="Arial", size=k.pop("size", 10), **k)
HEAD_FILL = PatternFill("solid", fgColor=BROWN)
TOTAL_FILL = PatternFill("solid", fgColor=CREAM)
_side = Side(style="thin", color="D9C9B2")
BORDER = Border(left=_side, right=_side, top=_side, bottom=_side)
MONEY, INT, PCT, DT = "#,##0", "#,##0", "0.0%", "dd/mm/yyyy hh:mm"


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = F(size=16, bold=True, color=BROWN)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = F(italic=True, color="7B6A5B")


def table(ws, row, headers, rows, fmts, total_cols=(), total_label="Tổng cộng"):
    """Ghi bảng bắt đầu ở hàng `row`. Trả về (hàng dữ liệu đầu, hàng dữ liệu cuối, hàng trống kế tiếp)."""
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font, cell.fill, cell.border = F(bold=True, color="FFFFFF"), HEAD_FILL, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    first = row + 1
    for r, values in enumerate(rows, first):
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.border = F(), BORDER
            if fmts[c - 1]:
                cell.number_format = fmts[c - 1]
    last = first + len(rows) - 1
    nxt = last + 1
    if total_cols and rows:
        ws.cell(row=nxt, column=1, value=total_label)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=nxt, column=c)
            cell.font, cell.fill, cell.border = F(bold=True), TOTAL_FILL, BORDER
            if c in total_cols:
                L = ws.cell(row=1, column=c).column_letter
                cell.value = f"=SUM({L}{first}:{L}{last})"
                cell.number_format = fmts[c - 1] or INT
        nxt += 1
    return first, last, nxt + 1


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w


def build_workbook(rep, categories=None):
    categories = categories or {}
    p, s, prev = rep["period"], rep["summary"], rep["prev"]
    now = datetime.now(VN).strftime("%d/%m/%Y %H:%M")
    wb = Workbook()

    # ---------- 1. Tổng quan ----------
    ws = wb.active
    ws.title = "Tổng quan"
    title(ws, "BÁO CÁO DOANH THU – BEAN & BREW CAFÉ", f"{p['label']}   ·   Xuất lúc {now} (giờ Việt Nam)")
    for c, h in enumerate(["Chỉ tiêu", p["label"], prev["label"], "Thay đổi so với kỳ trước"], 1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font, cell.fill, cell.border = F(bold=True, color="FFFFFF"), HEAD_FILL, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[4].height = 32
    rows = [
        ("Doanh thu (VNĐ)", s["revenue"], prev["revenue"], MONEY),
        ("Số đơn hợp lệ", s["orders"], prev["orders"], INT),
        ("Giá trị đơn trung bình (VNĐ)", "=IF(B6=0,0,B5/B6)", "=IF(C6=0,0,C5/C6)", MONEY),
    ]
    for i, (name, cur, pv, fmt) in enumerate(rows, 5):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=cur).number_format = fmt
        ws.cell(row=i, column=3, value=pv).number_format = fmt
        ws.cell(row=i, column=4, value=f'=IF(C{i}=0,"n/a",(B{i}-C{i})/C{i})').number_format = PCT
    extra = [("Đơn bị hủy", s["cancelled"]), ("Khách hàng (không trùng)", s["customers"]), ("Số món bán ra", s["itemsQty"])]
    for i, (name, val) in enumerate(extra, 8):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=val).number_format = INT
    for r in range(5, 11):
        for c in range(1, 5):
            cell = ws.cell(row=r, column=c)
            cell.font, cell.border = F(bold=(c == 1)), BORDER
            cell.alignment = Alignment(horizontal="left" if c == 1 else "right")
    notes = ["Ghi chú:",
             "• Doanh thu, số đơn, giá trị trung bình không tính đơn đã hủy; mốc thời gian là lúc khách đặt hàng (UTC+7).",
             "• Kỳ trước = tuần liền trước (hoặc tháng liền trước) để so sánh; “n/a” nghĩa là kỳ trước không có doanh thu.",
             "• Nguồn dữ liệu: MongoDB, collection orders. Các ô tổng và % thay đổi là công thức Excel."]
    for i, t in enumerate(notes, 12):
        ws.cell(row=i, column=1, value=t).font = F(bold=(i == 12), color="7B6A5B")
    widths(ws, [34, 26, 26, 24])

    # ---------- 2. Theo ngày ----------
    ws = wb.create_sheet("Theo ngày")
    title(ws, "Doanh thu theo ngày", p["label"])
    first, last, _ = table(ws, 4, ["Ngày", "Thứ", "Số đơn", "Doanh thu (VNĐ)"],
                           [(d["date"], d["weekday"], d["orders"], d["revenue"]) for d in rep["byDay"]],
                           [None, None, INT, MONEY], total_cols=(3, 4))
    widths(ws, [14, 12, 10, 18])
    ch = BarChart()
    ch.type, ch.title, ch.style = "col", "Doanh thu theo ngày (VNĐ)", 10
    ch.add_data(Reference(ws, min_col=4, min_row=4, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
    ch.legend, ch.height, ch.width = None, 8.5, 18
    ws.add_chart(ch, "F4")

    # ---------- 3. Phân tích ----------
    ws = wb.create_sheet("Phân tích")
    title(ws, "Phân tích doanh thu", p["label"])
    blocks = [
        ("Theo nhóm món", ["Nhóm", "Số lượng", "Doanh thu (VNĐ)"],
         [(categories.get(x["_id"], x["_id"]), x["qty"], x["revenue"]) for x in rep["byCategory"]]),
        ("Top 10 món bán chạy", ["Món", "Số lượng", "Doanh thu (VNĐ)"],
         [(f"{i}. {x['_id']}", x["qty"], x["revenue"]) for i, x in enumerate(rep["topItems"], 1)]),
        ("Theo hình thức phục vụ", ["Hình thức", "Số đơn", "Doanh thu (VNĐ)"],
         [(TY.get(x["_id"], x["_id"]), x["orders"], x["revenue"]) for x in rep["byType"]]),
        ("Theo phương thức thanh toán", ["Phương thức", "Số đơn", "Doanh thu (VNĐ)"],
         [(PAY.get(x["_id"], x["_id"]), x["orders"], x["revenue"]) for x in rep["byPayment"]]),
    ]
    row = 4
    for name, heads, data in blocks:
        ws.cell(row=row, column=1, value=name).font = F(size=12, bold=True, color=BROWN)
        _, _, row = table(ws, row + 1, heads, data, [None, INT, MONEY], total_cols=(2, 3))
        row += 0
    widths(ws, [36, 14, 20])

    # ---------- 4. Đơn hàng ----------
    ws = wb.create_sheet("Đơn hàng")
    orders = rep.get("orders", [])
    title(ws, "Chi tiết đơn hàng", f"{p['label']} · {len(orders)} đơn (kể cả đơn đã hủy)")
    heads = ["Mã đơn", "Thời gian", "Khách hàng", "SĐT", "Hình thức", "Bàn", "Món", "Tạm tính", "Phí giao", "Tổng (VNĐ)", "Thanh toán", "Trạng thái"]
    data = []
    for o in orders:
        items = "; ".join(f"{l['qty']}× {l['name']} ({l['variant']}" + (" + " + ", ".join(t["name"] for t in l["toppings"]) if l["toppings"] else "") + ")"
                          for l in o["items"])
        created = o["createdAt"].astimezone(VN).replace(tzinfo=None)
        data.append((o["code"], created, o["customer"]["name"], o["customer"]["phone"], TY.get(o["type"], o["type"]),
                     o.get("tableNumber"), items, o["subtotal"], o["shippingFee"], o["total"],
                     PAY.get(o["paymentMethod"], o["paymentMethod"]), ST.get(o["status"], o["status"])))
    first, last, nxt = table(ws, 4, heads, data, [None, DT, None, None, None, None, None, MONEY, MONEY, MONEY, None, None])
    if data:
        ws.cell(row=nxt - 1, column=9, value="Tổng doanh thu (không tính đơn hủy):").font = F(bold=True)
        cell = ws.cell(row=nxt - 1, column=10, value=f'=SUMIF(L{first}:L{last},"<>Đã hủy",J{first}:J{last})')
        cell.font, cell.number_format, cell.fill = F(bold=True), MONEY, TOTAL_FILL
        ws.auto_filter.ref = f"A4:L{last}"
    ws.freeze_panes = "A5"
    for r in range(first, last + 1):
        ws.cell(row=r, column=7).alignment = Alignment(wrap_text=True, vertical="top")
    widths(ws, [16, 17, 20, 13, 12, 6, 52, 12, 10, 14, 14, 14])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
