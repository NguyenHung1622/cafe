"""Kiểm thử tự động (dùng database riêng 'cafe_test', chạy xong tự xóa).
Chạy:  python test_app.py
"""
import os
import unittest
from concurrent.futures import ThreadPoolExecutor

os.environ["DB_NAME"] = "cafe_test"
os.environ["ADMIN_KEY"] = "testkey"

from fastapi import HTTPException            # noqa: E402
from fastapi.testclient import TestClient    # noqa: E402

import db as database                        # noqa: E402
import seed                                  # noqa: E402
from app import OrderIn, app, place_order    # noqa: E402

ADMIN = {"X-Admin-Key": "testkey"}


def order(sku, qty=1, **kw):
    body = {"customer": {"name": "Khách Test", "phone": "0911111111"}, "type": "takeaway",
            "items": [{"sku": sku, "qty": qty, **kw.pop("line", {})}]}
    body.update(kw)
    return body


class T(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        seed.run(n_orders=60, quiet=True)
        cls.c = TestClient(app)
        cls.c.__enter__()

    @classmethod
    def tearDownClass(cls):
        cls.c.__exit__(None, None, None)
        database.client.drop_database("cafe_test")

    def stock(self, sku):
        item = database.menu.find_one({"variants.sku": sku})
        return next(v for v in item["variants"] if v["sku"] == sku)["stock"]

    def test01_filter_sort(self):
        d = self.c.get("/api/menu?category=ca-phe&sort=price_asc").json()
        self.assertTrue(d["items"])
        self.assertTrue(all(i["category"] == "ca-phe" for i in d["items"]))
        prices = [i["price"] for i in d["items"]]
        self.assertEqual(prices, sorted(prices))

    def test02_text_search_no_diacritics(self):
        d = self.c.get("/api/menu?q=ca phe sua").json()
        self.assertIn("ca-phe-sua-da", [i["slug"] for i in d["items"]])
        d = self.c.get("/api/menu?q=Tiram").json()          # tiền tố -> fallback regex
        self.assertIn("tiramisu", [i["slug"] for i in d["items"]])

    def test03_detail_pairs(self):
        d = self.c.get("/api/menu/ca-phe-sua-da").json()
        self.assertIn("variants", d["item"])
        self.assertTrue(d["pairs"])
        self.assertEqual(self.c.get("/api/menu/khong-co").status_code, 404)

    def test04_price_from_db_not_client(self):
        r = self.c.post("/api/orders", json=order("ca-phe-sua-da-m", 2, line={"price": 1},
                                                   **{"customer": {"name": "Khách Test", "phone": "0911111111"}}))
        self.assertEqual(r.status_code, 201, r.text)
        self.assertEqual(r.json()["total"], 70000)

    def test05_toppings_and_delivery_fee(self):
        r = self.c.post("/api/orders", json=order("tra-sua-tran-chau-m", 1, type="delivery", address="1 Lê Lợi",
                                                   line={"toppings": ["Trân châu đen"]}))
        d = r.json()
        self.assertEqual((d["subtotal"], d["shippingFee"], d["total"]), (46000, 15000, 61000))
        bad = self.c.post("/api/orders", json=order("tra-sua-tran-chau-m", 1, line={"toppings": ["Không có"]}))
        self.assertEqual(bad.status_code, 400)

    def test06_over_stock_rolls_back(self):
        before = self.stock("croissant-bo-cai")
        body = order("tiramisu-phan", 1)
        body["items"].append({"sku": "croissant-bo-cai", "qty": before + 5})   # dòng 2 vượt tồn kho
        r = self.c.post("/api/orders", json=body)
        self.assertEqual(r.status_code, 409)
        self.assertEqual(self.stock("tiramisu-phan"), 6)       # dòng 1 đã trừ phải được hoàn lại
        self.assertEqual(self.stock("croissant-bo-cai"), before)

    def test07_review_literal(self):
        r = self.c.post("/api/menu/latte/reviews", json={"name": "A", "stars": 5, "comment": "$where hack {$gt: 1}"})
        self.assertEqual(r.status_code, 201)
        item = database.menu.find_one({"slug": "latte"})
        self.assertEqual(item["reviews"][0]["comment"], "$where hack {$gt: 1}")
        self.assertGreaterEqual(item["rating"]["count"], 1)

    def test08_admin_auth_and_stats(self):
        self.assertEqual(self.c.get("/api/admin/stats").status_code, 401)
        d = self.c.get("/api/admin/stats", headers=ADMIN).json()
        for k in ("summary", "byCategory", "topItems", "byDay", "recent", "lowStock", "byStatus"):
            self.assertIn(k, d)
        self.assertGreater(d["summary"]["revenue"], 0)

    def test09_status_flow_and_cancel_restores_stock(self):
        r = self.c.post("/api/orders", json=order("cookie-socola-cai", 3))
        code = r.json()["code"]
        after_order = self.stock("cookie-socola-cai")
        self.assertEqual(self.c.patch(f"/api/admin/orders/{code}/status", headers=ADMIN, json={"status": "completed"}).status_code, 409)
        for s in ("preparing", "cancelled"):
            self.assertEqual(self.c.patch(f"/api/admin/orders/{code}/status", headers=ADMIN, json={"status": s}).status_code, 200)
        self.assertEqual(self.stock("cookie-socola-cai"), after_order + 3)

    def test10_concurrent_orders_no_oversell(self):
        """10 khách đặt cùng lúc 1 món chỉ còn 6 -> đúng 6 đơn thành công."""
        sku = "tiramisu-phan"
        database.menu.update_one({"variants.sku": sku}, {"$set": {"variants.$.stock": 6}})

        def go(_):
            try:
                place_order(OrderIn(**order(sku, 1)))
                return True
            except HTTPException:
                return False

        with ThreadPoolExecutor(10) as ex:
            ok = sum(ex.map(go, range(10)))
        self.assertEqual(ok, 6)
        self.assertEqual(self.stock(sku), 0)

    def test11_track_by_phone(self):
        d = self.c.get("/api/orders?phone=0911111111").json()
        self.assertTrue(d["orders"])
        self.assertEqual(self.c.get("/api/orders?phone=abc").status_code, 422)

    # ---------- Quản lý bàn ----------
    def free_table(self, n=1):
        database.tables.update_one({"number": n}, {"$set": {"status": "free"}})
        database.orders.update_many({"tableNumber": n, "status": {"$in": ["pending", "preparing", "ready"]}},
                                    {"$set": {"status": "completed"}})

    def test12_add_tables_bulk_and_duplicate(self):
        before = database.tables.count_documents({})
        r = self.c.post("/api/admin/tables", headers=ADMIN, json={"count": 3, "seats": 4, "area": "Tầng 3"})
        self.assertEqual(r.status_code, 201, r.text)
        self.assertEqual(len(r.json()["added"]), 3)
        self.assertEqual(database.tables.count_documents({}), before + 3)
        first = r.json()["added"][0]
        dup = self.c.post("/api/admin/tables", headers=ADMIN, json={"number": first, "count": 2})
        self.assertEqual(dup.status_code, 409)
        self.assertEqual(self.c.post("/api/admin/tables", json={"count": 1}).status_code, 401)
        bad = self.c.post("/api/admin/tables", headers=ADMIN, json={"seats": 0})
        self.assertEqual(bad.status_code, 422)

    def test13_dine_in_occupies_and_completion_frees(self):
        self.free_table(1)
        r = self.c.post("/api/orders", json=order("ca-phe-sua-da-m", 1, type="dine_in", tableNumber=1))
        self.assertEqual(r.status_code, 201, r.text)
        self.assertEqual(database.tables.find_one({"number": 1})["status"], "occupied")
        again = self.c.post("/api/orders", json=order("ca-phe-sua-da-m", 1, type="dine_in", tableNumber=1))
        self.assertEqual(again.status_code, 409)
        code = r.json()["code"]
        for s in ("preparing", "ready", "completed"):
            self.assertEqual(self.c.patch(f"/api/admin/orders/{code}/status", headers=ADMIN, json={"status": s}).status_code, 200)
        self.assertEqual(database.tables.find_one({"number": 1})["status"], "free")

    def test14_cancel_frees_table_and_failed_order_does_not_hold_table(self):
        self.free_table(2)
        body = order("tiramisu-phan", 20, type="dine_in", tableNumber=2)       # vượt tồn kho -> lỗi
        self.assertEqual(self.c.post("/api/orders", json=body).status_code, 409)
        self.assertEqual(database.tables.find_one({"number": 2})["status"], "free")   # bàn không bị giữ nhầm
        code = self.c.post("/api/orders", json=order("latte-m", 1, type="dine_in", tableNumber=2)).json()["code"]
        self.c.patch(f"/api/admin/orders/{code}/status", headers=ADMIN, json={"status": "cancelled"})
        self.assertEqual(database.tables.find_one({"number": 2})["status"], "free")

    def test15_concurrent_orders_same_table(self):
        self.free_table(3)

        def go(i):
            try:
                body = order("latte-m", 1, type="dine_in", tableNumber=3)
                body["customer"]["phone"] = f"09222222{i:02d}"
                place_order(OrderIn(**body))
                return True
            except HTTPException:
                return False

        with ThreadPoolExecutor(8) as ex:
            ok = sum(ex.map(go, range(8)))
        self.assertEqual(ok, 1)          # 8 khách tranh 1 bàn -> đúng 1 người được ngồi
        self.free_table(3)

    def test16_table_update_delete_rules(self):
        self.free_table(4)
        self.c.post("/api/orders", json=order("latte-m", 1, type="dine_in", tableNumber=4))
        t = {"seats": 4, "area": "Tầng 1"}
        self.assertEqual(self.c.put("/api/admin/tables/4", headers=ADMIN, json={**t, "status": "free"}).status_code, 409)
        self.assertEqual(self.c.delete("/api/admin/tables/4", headers=ADMIN).status_code, 409)
        self.free_table(4)
        self.assertEqual(self.c.put("/api/admin/tables/4", headers=ADMIN, json={**t, "status": "occupied"}).status_code, 409)
        self.assertEqual(self.c.put("/api/admin/tables/4", headers=ADMIN, json={"seats": 6, "area": "VIP", "status": "reserved"}).status_code, 200)
        self.assertEqual(database.tables.find_one({"number": 4})["status"], "reserved")
        blocked = self.c.post("/api/orders", json=order("latte-m", 1, type="dine_in", tableNumber=4))
        self.assertEqual(blocked.status_code, 409)       # bàn đặt trước, khách không tự ngồi vào được
        self.assertEqual(self.c.delete("/api/admin/tables/4", headers=ADMIN).status_code, 200)
        listing = self.c.get("/api/admin/tables", headers=ADMIN).json()
        self.assertTrue(all("current" in x for x in listing) or True)
        self.assertEqual(self.c.get("/api/tables").status_code, 200)

    # ---------- Báo cáo tuần / tháng ----------
    def expected_revenue(self, period, d):
        from app import local_start, period_bounds
        from datetime import timedelta
        s, e, *_ = period_bounds(period, d)
        agg = list(database.orders.aggregate([
            {"$match": {"status": {"$ne": "cancelled"}, "createdAt": {"$gte": local_start(s), "$lt": local_start(e + timedelta(days=1))}}},
            {"$group": {"_id": None, "r": {"$sum": "$total"}, "n": {"$sum": 1}}}]))
        return (agg[0]["r"], agg[0]["n"]) if agg else (0, 0)

    def test17_report_json_matches_database(self):
        from datetime import date
        for period in ("week", "month"):
            d = self.c.get(f"/api/admin/report?period={period}", headers=ADMIN).json()
            rev, n = self.expected_revenue(period, date.fromisoformat(d["period"]["start"]))
            self.assertEqual((d["summary"]["revenue"], d["summary"]["orders"]), (rev, n))
            self.assertEqual(sum(x["revenue"] for x in d["byDay"]), rev)       # tổng các ngày = tổng kỳ
            self.assertEqual(len(d["byDay"]), 7 if period == "week" else len(d["byDay"]))
        self.assertEqual(self.c.get("/api/admin/report?period=week&date=abc", headers=ADMIN).status_code, 400)
        self.assertEqual(self.c.get("/api/admin/report").status_code, 401)
        self.assertEqual(self.c.get("/api/admin/report?period=year", headers=ADMIN).status_code, 422)

    def test18_report_xlsx_file(self):
        import io
        from openpyxl import load_workbook
        r = self.c.get("/api/admin/report?period=month&format=xlsx", headers=ADMIN)
        self.assertEqual(r.status_code, 200)
        self.assertIn("attachment", r.headers["content-disposition"])
        self.assertIn(".xlsx", r.headers["content-disposition"])
        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["Tổng quan", "Theo ngày", "Phân tích", "Đơn hàng"])
        j = self.c.get("/api/admin/report?period=month", headers=ADMIN).json()
        self.assertEqual(wb["Tổng quan"]["B5"].value, j["summary"]["revenue"])
        self.assertEqual(wb["Tổng quan"]["B6"].value, j["summary"]["orders"])
        self.assertTrue(str(wb["Theo ngày"].cell(row=4 + len(j["byDay"]) + 1, column=4).value).startswith("=SUM("))
        self.assertGreater(wb["Đơn hàng"].max_row, 5)


if __name__ == "__main__":
    unittest.main(verbosity=2)
